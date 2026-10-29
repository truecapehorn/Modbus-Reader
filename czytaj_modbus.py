import asyncio
import argparse
import json
import time
from datetime import datetime
from pathlib import Path

try:
    from pymodbus.client import ModbusSerialClient
    from pymodbus.exceptions import ModbusException
except ImportError as error:
    raise SystemExit(
        "Brak biblioteki pymodbus. Zainstaluj zaleznosci poleceniem: "
        "python -m pip install -r requirements.txt"
    ) from error

try:
    from rich.console import Console, Group
    from rich.table import Table
    from rich.text import Text
    from textual.app import App, ComposeResult
    from textual.widgets import DataTable, Footer, Header, Static
except ImportError as error:
    raise SystemExit(
        "Brak wymaganej biblioteki Rich lub Textual. Zainstaluj zaleznosci "
        "poleceniem: python -m pip install -r requirements.txt"
    ) from error


def load_register_info(path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise SystemExit(
            "Brak biblioteki openpyxl. Zainstaluj ja poleceniem: "
            "python -m pip install -r requirements.txt"
        ) from error

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise SystemExit(f"Nie znaleziono pliku z mapa zmiennych: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        sheet = workbook["Status_var"]
    except KeyError as error:
        workbook.close()
        raise SystemExit(
            f"Arkusz 'Status_var' nie istnieje w pliku {workbook_path}"
        ) from error

    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        workbook.close()
        raise SystemExit(f"Arkusz 'Status_var' w {workbook_path} jest pusty.")

    columns = {str(name).strip(): index for index, name in enumerate(headers) if name}
    required = {"Address", "Name", "IEC type"}
    missing = required - columns.keys()
    if missing:
        workbook.close()
        raise SystemExit(
            "Brak wymaganych kolumn w arkuszu Status_var: "
            + ", ".join(sorted(missing))
        )

    metadata_columns = {
        "name": "Name",
        "description": "Description",
        "iec_type": "IEC type",
        "min": "Min",
        "max": "Max",
        "scale": "Scale",
        "offset": "Offset",
        "unit": "Unit",
    }

    def cell_text(row, column_name):
        index = columns.get(column_name)
        if index is None or index >= len(row) or row[index] is None:
            return "-"
        return str(row[index]).strip() or "-"

    register_info = {}
    for row_number, row in enumerate(rows, start=2):
        if not row or row[columns["Address"]] is None:
            continue
        try:
            address = int(row[columns["Address"]])
        except (TypeError, ValueError) as error:
            workbook.close()
            raise SystemExit(
                f"Nieprawidlowy adres w arkuszu Status_var, wiersz {row_number}."
            ) from error

        register_info[address] = {
            key: cell_text(row, column_name)
            for key, column_name in metadata_columns.items()
        }

    workbook.close()
    if not register_info:
        raise SystemExit(f"Nie znaleziono adresow w arkuszu Status_var: {workbook_path}")
    return register_info


def contiguous_ranges(addresses):
    sorted_addresses = sorted(set(addresses))
    if not sorted_addresses:
        return

    range_start = previous = sorted_addresses[0]
    for address in sorted_addresses[1:]:
        if address != previous + 1:
            yield range_start, previous - range_start + 1
            range_start = address
        previous = address
    yield range_start, previous - range_start + 1


def read_register_range(client, slave, start, count):
    try:
        if not client.connected and not client.connect():
            raise ConnectionError("Nie udalo sie polaczyc z portem szeregowym.")
        response = client.read_holding_registers(
            start, count=count, slave=slave
        )
    except (ModbusException, OSError) as error:
        return {}, {
            address: str(error) or error.__class__.__name__
            for address in range(start, start + count)
        }

    if response.isError():
        if getattr(response, "exception_code", None) == 2 and count > 1:
            left_count = count // 2
            left_values, left_errors = read_register_range(
                client, slave, start, left_count
            )
            right_values, right_errors = read_register_range(
                client, slave, start + left_count, count - left_count
            )
            return left_values | right_values, left_errors | right_errors

        error = str(response)
        return {}, {
            address: error
            for address in range(start, start + count)
        }

    values = {
        address: str(value)
        for address, value in enumerate(response.registers, start=start)
    }
    if len(values) != count:
        error = (
            f"Odpowiedz zawiera {len(values)} rejestrow; "
            f"oczekiwano {count}."
        )
        return values, {
            address: error
            for address in range(start, start + count)
            if address not in values
        }
    return values, {}


def read_configured_registers(
    client, slave, register_info, start, count, on_update=None
):
    last_address = start + count - 1
    addresses = []
    for address, info in register_info.items():
        if start <= address <= last_address:
            addresses.append(address)
            if info["iec_type"] == "UDINT" and address < last_address:
                addresses.append(address + 1)

    values = {}
    errors = {}
    for block_start, block_count in contiguous_ranges(addresses):
        block_values, block_errors = read_register_range(
            client, slave, block_start, block_count
        )
        values.update(block_values)
        errors.update(block_errors)
        if on_update is not None:
            on_update(block_values, block_errors)
    return values, errors


def build_dashboard(start, count, values, register_info, errors=None):
    errors = errors or {}
    table = Table(title="Odczyt rejestrow Modbus RTU", expand=True)
    for column in (
        "Adres",
        "Name",
        "Description",
        "IEC type",
        "Min",
        "Max",
        "Scale",
        "Offset",
        "Unit",
        "Wartosc",
    ):
        table.add_column(column, no_wrap=column in {"Adres", "IEC type"})

    failed_reads = 0
    for address in range(start, start + count):
        info = register_info.get(address)
        if info is None:
            continue

        value = values.get(address, "Brak odczytu")
        if address in errors:
            value = f"BLAD: {errors[address].splitlines()[-1]}"
            value = value.replace("\n", " ")
            failed_reads += 1
        if info["iec_type"] == "BOOL" and value in ("0", "1"):
            value = f"{value} ({'FALSE' if value == '0' else 'TRUE'})"

        if value.endswith("(TRUE)"):
            value_style = "bold green"
        elif value.endswith("(FALSE)"):
            value_style = "bold red"
        elif value.startswith("BLAD:"):
            value_style = "bold white on red"
        elif value == "Brak odczytu":
            value_style = "dim"
        elif value.lstrip("-").isdigit():
            value_style = "bold cyan"
        else:
            value_style = ""

        table.add_row(
            str(address),
            info["name"],
            info["description"],
            info["iec_type"],
            info["min"],
            info["max"],
            info["scale"],
            info["offset"],
            info["unit"],
            value,
            style=value_style,
        )

    status = Text(
        f"{datetime.now():%Y-%m-%d %H:%M:%S} | "
        f"zakres {start}-{start + count - 1} | "
        f"odczyty bledne: {failed_reads}"
    )
    return Group(status, table)


def styled_value(value, iec_type):
    if iec_type == "BOOL" and value in ("0", "1"):
        state = "TRUE" if value == "1" else "FALSE"
        style = "bold green" if value == "1" else "bold red"
        return Text(f"{value} ({state})", style=style)
    if value.startswith("BLAD:"):
        return Text(value, style="bold white on red")
    if value == "Brak odczytu":
        return Text(value, style="dim")
    if value.lstrip("-").isdigit():
        return Text(value, style="bold cyan")
    return Text(value)


class RegisterMonitor(App):
    CSS = """
    #status {
        height: 2;
        padding: 0 1;
    }
    DataTable {
        height: 1fr;
    }
    """
    BINDINGS = [
        ("q", "quit", "Koniec"),
        ("home", "scroll_top", "Poczatek tabeli"),
        ("end", "scroll_bottom", "Koniec tabeli"),
    ]

    def __init__(
        self,
        client,
        slave,
        register_info,
        start,
        count,
        interval,
        no_color=False,
    ):
        super().__init__(ansi_color=not no_color)
        self.client = client
        self.slave = slave
        self.register_info = register_info
        self.start = start
        self.count = count
        self.interval = interval
        self.value_column = None

    def compose(self) -> ComposeResult:
        yield Header(show_clock=False)
        yield Static("Ladowanie odczytow...", id="status")
        yield DataTable(id="registers", zebra_stripes=True)
        yield Footer()

    def on_mount(self):
        table = self.query_one("#registers", DataTable)
        for label, key in (
            ("Adres", "address"),
            ("Name", "name"),
            ("Description", "description"),
            ("IEC type", "iec_type"),
            ("Min", "min"),
            ("Max", "max"),
            ("Scale", "scale"),
            ("Offset", "offset"),
            ("Unit", "unit"),
            ("Wartosc", "value"),
        ):
            column_key = table.add_column(label, key=key)
            if key == "value":
                self.value_column = column_key

        for address, info in sorted(self.register_info.items()):
            if not self.start <= address < self.start + self.count:
                continue
            table.add_row(
                str(address),
                info["name"],
                info["description"],
                info["iec_type"],
                info["min"],
                info["max"],
                info["scale"],
                info["offset"],
                info["unit"],
                Text("Oczekiwanie...", style="dim"),
                key=str(address),
            )

        self.run_worker(self.poll_loop(), name="modbus-poll", exclusive=True)

    async def poll_loop(self):
        while True:
            values, errors = await asyncio.to_thread(
                read_configured_registers,
                self.client,
                self.slave,
                self.register_info,
                self.start,
                self.count,
            )

            table = self.query_one("#registers", DataTable)
            for address in sorted(self.register_info):
                if not self.start <= address < self.start + self.count:
                    continue
                info = self.register_info[address]
                value = values.get(address, "Brak odczytu")
                if address in errors:
                    value = f"BLAD: {errors[address].splitlines()[-1]}"
                table.update_cell(
                    str(address),
                    self.value_column,
                    styled_value(value, info["iec_type"]),
                )

            self.query_one("#status", Static).update(
                f"{datetime.now():%Y-%m-%d %H:%M:%S} | "
                f"{len(self.register_info)} zmiennych | "
                f"bledne odczyty: {len(errors)} | "
                "Przewijaj tabela; q = wyjscie"
            )
            await asyncio.sleep(self.interval)

    def action_scroll_top(self):
        self.query_one("#registers", DataTable).scroll_home(animate=False)

    def action_scroll_bottom(self):
        self.query_one("#registers", DataTable).scroll_end(animate=False)


def remove_json_line_comments(contents):
    cleaned = []
    in_string = False
    escaped = False
    index = 0

    while index < len(contents):
        char = contents[index]
        if in_string:
            cleaned.append(char)
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            index += 1
        elif char == '"':
            in_string = True
            cleaned.append(char)
            index += 1
        elif char == "/" and index + 1 < len(contents) and contents[index + 1] == "/":
            while index < len(contents) and contents[index] != "\n":
                cleaned.append(" ")
                index += 1
        else:
            cleaned.append(char)
            index += 1

    return "".join(cleaned)


def load_config(path):
    config_path = Path(path)
    try:
        with config_path.open(encoding="utf-8") as config_file:
            config = json.loads(remove_json_line_comments(config_file.read()))
    except OSError as error:
        raise SystemExit(
            f"Nie mozna odczytac pliku konfiguracji {config_path}: {error}"
        ) from error
    except json.JSONDecodeError as error:
        raise SystemExit(
            f"Nieprawidlowy JSON w pliku konfiguracji {config_path}, "
            f"wiersz {error.lineno}, kolumna {error.colno}."
        ) from error

    if not isinstance(config, dict):
        raise SystemExit(f"Plik konfiguracji {config_path} musi zawierac obiekt JSON.")

    expected_types = {
        "port": str,
        "slave": int,
        "baud": int,
        "start": int,
        "count": int,
        "interval": float,
        "timeout": float,
        "parity": str,
        "bytesize": int,
        "stopbits": float,
        "metadata": str,
        "no_clear": bool,
        "no_color": bool,
    }
    unknown = config.keys() - expected_types.keys()
    missing = expected_types.keys() - config.keys()
    if unknown:
        raise SystemExit(
            f"Nieznane ustawienia w pliku {config_path}: "
            + ", ".join(sorted(unknown))
        )
    if missing:
        raise SystemExit(
            f"Brak ustawien w pliku {config_path}: "
            + ", ".join(sorted(missing))
        )

    for key, expected_type in expected_types.items():
        value = config[key]
        valid = (
            type(value) is expected_type
            or expected_type is float and type(value) is int
        )
        if not valid:
            raise SystemExit(
                f"Ustawienie '{key}' w pliku {config_path} ma nieprawidlowy typ."
            )
        if expected_type is float:
            config[key] = float(value)

    if config["parity"] not in {"N", "E", "O"}:
        raise SystemExit(
            f"Ustawienie 'parity' w pliku {config_path} musi miec wartosc N, E lub O."
        )
    if config["bytesize"] not in {5, 6, 7, 8}:
        raise SystemExit(
            f"Ustawienie 'bytesize' w pliku {config_path} musi miec wartosc 5, 6, 7 lub 8."
        )
    if config["stopbits"] not in {1.0, 1.5, 2.0}:
        raise SystemExit(
            f"Ustawienie 'stopbits' w pliku {config_path} musi miec wartosc 1, 1.5 lub 2."
        )

    metadata_path = Path(config["metadata"])
    if not metadata_path.is_absolute():
        config["metadata"] = str(config_path.resolve().parent / metadata_path)
    return config


def main():
    config_parser = argparse.ArgumentParser(add_help=False)
    config_parser.add_argument(
        "--config",
        default=str(Path(__file__).parent / "config.json"),
        help="Plik JSON z domyslnymi ustawieniami (obsluguje komentarze //)",
    )
    config_args, _ = config_parser.parse_known_args()
    config = load_config(config_args.config)

    parser = argparse.ArgumentParser(
        description="Cykliczny odczyt rejestrow Modbus RTU przez pymodbus"
    )
    parser.add_argument(
        "--config",
        default=config_args.config,
        help="Plik JSON z domyslnymi ustawieniami (obsluguje komentarze //)",
    )
    parser.add_argument("--port", help="Port, np. /dev/ttyUSB0 lub COM3")
    parser.add_argument("--slave", type=int, help="Adres urzadzenia Modbus")
    parser.add_argument("--baud", type=int, help="Predkosc transmisji")
    parser.add_argument("--start", type=int, help="Pierwszy adres rejestru")
    parser.add_argument("--count", type=int, help="Liczba rejestrow")
    parser.add_argument(
        "--interval",
        type=float,
        help="Przerwa miedzy odczytami w sekundach",
    )
    parser.add_argument(
        "--timeout",
        type=float,
        help="Limit czasu pojedynczego odczytu w sekundach",
    )
    parser.add_argument(
        "--parity",
        choices=("N", "E", "O"),
        help="Parzystosc transmisji",
    )
    parser.add_argument(
        "--bytesize",
        type=int,
        choices=(5, 6, 7, 8),
        help="Liczba bitow danych",
    )
    parser.add_argument(
        "--stopbits",
        type=float,
        choices=(1, 1.5, 2),
        help="Liczba bitow stopu",
    )
    clear_group = parser.add_mutually_exclusive_group()
    clear_group.add_argument(
        "--no-clear",
        action="store_true",
        help="Nie czysc terminala; dopisuj kolejne odczyty (przydatne do logowania)",
    )
    clear_group.add_argument(
        "--clear",
        dest="no_clear",
        action="store_false",
        help="Uzyj interaktywnej tabeli",
    )
    color_group = parser.add_mutually_exclusive_group()
    color_group.add_argument(
        "--no-color",
        action="store_true",
        help="Wylacz kolorowanie wartosci",
    )
    color_group.add_argument(
        "--color",
        dest="no_color",
        action="store_false",
        help="Wlacz kolorowanie wartosci",
    )
    parser.add_argument(
        "--metadata",
        help="Plik XLSX z arkuszem Status_var",
    )
    parser.set_defaults(**config)
    args = parser.parse_args()

    register_info = load_register_info(args.metadata)
    console = Console(no_color=args.no_color)

    configured_addresses = [
        address
        for address in register_info
        if args.start <= address < args.start + args.count
    ]
    if not configured_addresses:
        raise SystemExit("Zakres nie zawiera zadnego adresu z arkusza Status_var.")

    client = ModbusSerialClient(
        port=args.port,
        baudrate=args.baud,
        parity=args.parity,
        bytesize=args.bytesize,
        stopbits=args.stopbits,
        timeout=args.timeout,
    )
    try:
        if args.no_clear:
            while True:
                values, errors = read_configured_registers(
                    client,
                    args.slave,
                    register_info,
                    args.start,
                    args.count,
                )
                console.print(
                    build_dashboard(
                        args.start,
                        args.count,
                        values,
                        register_info,
                        errors,
                    )
                )
                time.sleep(args.interval)
        else:
            RegisterMonitor(
                client=client,
                slave=args.slave,
                register_info=register_info,
                start=args.start,
                count=args.count,
                interval=args.interval,
                no_color=args.no_color,
            ).run()

    except KeyboardInterrupt:
        console.print("\nZatrzymano odczyt.")
    finally:
        client.close()


if __name__ == "__main__":
    main()
